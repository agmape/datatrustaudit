import os
import json
import re
import io
import tempfile
from fastapi import FastAPI, Request
from fastapi.exceptions import RequestValidationError
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
import requests
from bs4 import BeautifulSoup
from datetime import datetime
from typing import Dict, List, Any, Optional
from urllib.parse import urlparse, urljoin
from contextlib import asynccontextmanager

# Load environment variables — override=True ensures .env always wins over system env
from dotenv import load_dotenv
load_dotenv(override=True)

# Google Gemini AI - Nova biblioteca
from google import genai
from google.genai import types

# Configure Gemini
GOOGLE_API_KEY = os.getenv("GOOGLE_API_KEY")
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-2.0-flash-lite")

# Cliente Gemini
gemini_client = None
if GOOGLE_API_KEY:
    gemini_client = genai.Client(api_key=GOOGLE_API_KEY)
    print(f"[OK] Gemini AI configured with model: {GEMINI_MODEL}")
else:
    print("[WARN] GOOGLE_API_KEY not configured. Chat will use rule-based responses.")

# Database path
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
STATIC_DIR = os.path.join(BASE_DIR, "static", "dist")  # Vite builds to dist folder
# Storage and DB
from db.database import engine, get_db, sync_sqlite_schema
from db.models import Base, User, Scan, Payment
from api import auth, payments, scans, audit

DB_PATH = os.path.join(BASE_DIR, "db", "tags.json")

@asynccontextmanager
async def lifespan(app: FastAPI):
    # Startup
    print(" DataTrust Audit API starting...")
    try:
        Base.metadata.create_all(bind=engine)
        sync_sqlite_schema()
        print(" Database tables created via SQLAlchemy")
    except Exception as e:
        print(f" Could not sync DB on startup: {e}")
        
    yield
    # Shutdown
    print(" DataTrust Audit API shutting down...")


app = FastAPI(
    title="DataTrust Audit API",
    description="Auditoria de Google Tag Manager e Conformidade LGPD",
    version="2.0.0",
    lifespan=lifespan
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.exception_handler(RequestValidationError)
async def validation_exception_handler(request: Request, exc: RequestValidationError):
    import traceback
    if import_meta_dev():
        traceback.print_exc()
    return JSONResponse(
        {
            "success": False,
            "error": {
                "code": "VALIDATION_ERROR",
                "message": "Dados inválidos.",
                "user_message": "Verifique seus dados e tente novamente.",
                "path": str(request.url.path),
                "method": request.method,
                "details": exc.errors(),
            },
        },
        status_code=422,
    )


@app.exception_handler(Exception)
async def global_exception_handler(request: Request, exc: Exception):
    import traceback
    # Always print full traceback to backend terminal — never hide errors
    traceback.print_exc()
    return JSONResponse(
        {
            "success": False,
            "error": {
                "code": "INTERNAL_SERVER_ERROR",
                "message": "Erro interno da API.",
                "user_message": (
                    "A API encontrou um erro interno. "
                    "Verifique os logs do servidor para identificar a causa real."
                ),
                "path": str(request.url.path),
                "method": request.method,
                "exception": type(exc).__name__,
                "debug_hint": "Veja o traceback no terminal do backend.",
            },
        },
        status_code=500,
    )


def import_meta_dev():
    """Check if we're in development mode."""
    return os.getenv("ENVIRONMENT", "development") == "development"


# ── Health check ──
@app.get("/health")
async def health_check():
    """Health check endpoint — verifies DB, env vars, and API status."""
    checks = {
        "status": "ok",
        "database": "unknown",
        "google_api_key": bool(GOOGLE_API_KEY),
        "supabase_jwt_secret": bool(os.getenv("SUPABASE_JWT_SECRET", "")),
        "admin_email": os.getenv("ADMIN_TEST_EMAIL", "not set"),
        "environment": os.getenv("ENVIRONMENT", "not set"),
    }
    try:
        from db.database import SessionLocal
        from sqlalchemy import text as sa_text
        db = SessionLocal()
        db.execute(sa_text("SELECT 1"))
        db.close()
        checks["database"] = "ok"
    except Exception as e:
        checks["database"] = f"error: {str(e)[:100]}"
        checks["status"] = "degraded"
    return checks


# Registra rotas da API
app.include_router(auth.router)
app.include_router(payments.router)
app.include_router(scans.router)
app.include_router(audit.router)


# Frontend static delivery.
#
# On Vercel, app.frontend() is the supported FastAPI integration and lets
# Vercel promote the generated Vite files to its static CDN. Locally (where
# plain FastAPI may not expose app.frontend), fall back to StaticFiles.
#
# IMPORTANT: never instantiate StaticFiles for a directory that does not exist.
# Doing that at import time crashes the whole serverless function with
# FUNCTION_INVOCATION_FAILED before /health can even respond.
FRONTEND_INDEX = os.path.join(STATIC_DIR, "index.html")
FRONTEND_ASSETS = os.path.join(STATIC_DIR, "assets")
FRONTEND_AVAILABLE = os.path.isfile(FRONTEND_INDEX)
VERCEL_FRONTEND_REGISTERED = False

if FRONTEND_AVAILABLE:
    _frontend = getattr(app, "frontend", None)
    if callable(_frontend):
        _frontend("/", directory=STATIC_DIR)
        VERCEL_FRONTEND_REGISTERED = True
        print(f"[OK] Vercel frontend registered from {STATIC_DIR}")
    elif os.path.isdir(FRONTEND_ASSETS):
        app.mount("/assets", StaticFiles(directory=FRONTEND_ASSETS), name="assets")
        print(f"[OK] Static assets mounted from {FRONTEND_ASSETS}")
else:
    print(
        f"[WARN] Frontend build not found at {STATIC_DIR}. "
        "API will still start; run the Vite production build."
    )


@app.get("/")
async def serve_index():
    """Serve the SPA root without making application startup depend on it."""
    if not os.path.isfile(FRONTEND_INDEX):
        return JSONResponse(
            {
                "status": "ok",
                "service": "DataTrust Audit API",
                "frontend": "build_not_found",
                "health": "/health",
            },
            status_code=200,
        )
    return FileResponse(
        FRONTEND_INDEX,
        headers={"Cache-Control": "no-cache, no-store, must-revalidate", "Pragma": "no-cache"},
    )

@app.get("/favicon.ico")
async def serve_favicon():
    """Serve o favicon"""
    favicon_path = os.path.join(STATIC_DIR, "favicon.ico")
    if os.path.exists(favicon_path):
        return FileResponse(favicon_path)
    return JSONResponse({"error": "not found"}, status_code=404)


@app.get("/robots.txt")
async def serve_robots():
    """Serve robots.txt"""
    robots_path = os.path.join(STATIC_DIR, "robots.txt")
    if os.path.exists(robots_path):
        return FileResponse(robots_path)
    return JSONResponse({"error": "not found"}, status_code=404)


# ============================================================
# TAG DETECTION PATTERNS - Real patterns for real detection
# ============================================================

TAG_PATTERNS = {
    "gtm": {
        "name": "Google Tag Manager",
        "patterns": [
            r"googletagmanager\.com/gtm\.js",
            r"GTM-[A-Z0-9]{6,8}",
            r"google_tag_manager"
        ],
        "type": "tag_manager",
        "data_collected": ["Eventos personalizados", "Dados de e-commerce", "Interações do usuário"],
        "lgpd_risk": "medium"
    },
    "ga4": {
        "name": "Google Analytics 4",
        "patterns": [
            r"gtag\s*\(\s*['\"]config['\"]",
            r"G-[A-Z0-9]{10,12}",
            r"google-analytics\.com/g/collect",
            r"googletagmanager\.com/gtag/js"
        ],
        "type": "analytics",
        "data_collected": ["Endereço IP", "Comportamento de navegação", "Localização geográfica", "Dados demográficos"],
        "lgpd_risk": "high"
    },
    "ua": {
        "name": "Universal Analytics (Obsoleto)",
        "patterns": [
            r"UA-\d{6,10}-\d{1,3}",
            r"google-analytics\.com/analytics\.js",
            r"ga\s*\(\s*['\"]create['\"]"
        ],
        "type": "analytics",
        "data_collected": ["Endereço IP", "Comportamento de navegação", "Cookies persistentes"],
        "lgpd_risk": "critical"
    },
    "meta_pixel": {
        "name": "Meta Pixel (Facebook)",
        "patterns": [
            r"connect\.facebook\.net.*fbevents\.js",
            r"fbq\s*\(\s*['\"]init['\"]",
            r"\d{15,16}"  # Facebook Pixel ID
        ],
        "type": "advertising",
        "data_collected": ["Fingerprint do navegador", "Comportamento de compra", "Remarketing cross-site"],
        "lgpd_risk": "critical"
    },
    "google_ads": {
        "name": "Google Ads Conversion",
        "patterns": [
            r"googleads\.g\.doubleclick\.net",
            r"AW-\d{9,11}",
            r"gtag\s*\(\s*['\"]event['\"],\s*['\"]conversion['\"]"
        ],
        "type": "advertising",
        "data_collected": ["Conversões", "Dados de remarketing", "Audiências personalizadas"],
        "lgpd_risk": "high"
    },
    "hotjar": {
        "name": "Hotjar",
        "patterns": [
            r"static\.hotjar\.com",
            r"hj\s*\(\s*['\"]trigger['\"]",
            r"hjid"
        ],
        "type": "heatmap",
        "data_collected": ["Gravações de sessão", "Mapas de calor", "Dados de formulário"],
        "lgpd_risk": "critical"
    },
    "clarity": {
        "name": "Microsoft Clarity",
        "patterns": [
            r"clarity\.ms",
            r"clarity\s*\("
        ],
        "type": "heatmap",
        "data_collected": ["Gravações de sessão", "Mapas de clique", "Scroll depth"],
        "lgpd_risk": "high"
    },
    "tiktok_pixel": {
        "name": "TikTok Pixel",
        "patterns": [
            r"analytics\.tiktok\.com",
            r"ttq\.track",
            r"ttq\.identify"
        ],
        "type": "advertising",
        "data_collected": ["Eventos de conversão", "Remarketing", "Audiências"],
        "lgpd_risk": "critical"
    },
    "linkedin_insight": {
        "name": "LinkedIn Insight Tag",
        "patterns": [
            r"snap\.licdn\.com",
            r"linkedin\.com/px",
            r"_linkedin_partner_id"
        ],
        "type": "advertising",
        "data_collected": ["Conversões B2B", "Remarketing profissional"],
        "lgpd_risk": "high"
    },
    "intercom": {
        "name": "Intercom",
        "patterns": [
            r"widget\.intercom\.io",
            r"Intercom\s*\(",
            r"intercomSettings"
        ],
        "type": "support",
        "data_collected": ["Mensagens de chat", "Email", "Nome do usuário", "Dados de navegação"],
        "lgpd_risk": "high"
    },
    "zendesk": {
        "name": "Zendesk",
        "patterns": [
            r"static\.zdassets\.com",
            r"zE\s*\(",
            r"zendesk"
        ],
        "type": "support",
        "data_collected": ["Tickets de suporte", "Email", "Histórico de chat"],
        "lgpd_risk": "medium"
    },
    "cookiebot": {
        "name": "Cookiebot (CMP)",
        "patterns": [
            r"consent\.cookiebot\.com",
            r"Cookiebot",
            r"CookieConsent"
        ],
        "type": "consent",
        "data_collected": ["Preferências de consentimento"],
        "lgpd_risk": "none"
    },
    "onetrust": {
        "name": "OneTrust (CMP)",
        "patterns": [
            r"cdn\.cookielaw\.org",
            r"OptanonWrapper",
            r"OneTrust"
        ],
        "type": "consent",
        "data_collected": ["Preferências de consentimento"],
        "lgpd_risk": "none"
    },
    "didomi": {
        "name": "Didomi (CMP)",
        "patterns": [
            r"sdk\.privacy-center\.org",
            r"Didomi",
            r"didomi"
        ],
        "type": "consent",
        "data_collected": ["Preferências de consentimento"],
        "lgpd_risk": "none"
    },
    # ============ NOVAS TAGS ADICIONADAS ============
    "pinterest": {
        "name": "Pinterest Tag",
        "patterns": [
            r"pintrk\s*\(",
            r"s\.pinimg\.com/ct/core\.js",
            r"pinterest\.com/ct"
        ],
        "type": "advertising",
        "data_collected": ["Conversões", "Remarketing", "Audiências de interesse"],
        "lgpd_risk": "high"
    },
    "twitter_pixel": {
        "name": "Twitter/X Pixel",
        "patterns": [
            r"static\.ads-twitter\.com",
            r"twq\s*\(",
            r"twitter\.com/i/adsct"
        ],
        "type": "advertising",
        "data_collected": ["Conversões", "Tailored Audiences", "Remarketing"],
        "lgpd_risk": "high"
    },
    "snapchat_pixel": {
        "name": "Snapchat Pixel",
        "patterns": [
            r"sc-static\.net/scevent\.min\.js",
            r"snaptr\s*\(",
            r"tr\.snapchat\.com"
        ],
        "type": "advertising",
        "data_collected": ["Conversões", "Remarketing", "Snapchat Audiences"],
        "lgpd_risk": "high"
    },
    "segment": {
        "name": "Segment",
        "patterns": [
            r"cdn\.segment\.com",
            r"analytics\.js",
            r"analytics\.identify",
            r"analytics\.track"
        ],
        "type": "analytics",
        "data_collected": ["Eventos personalizados", "Identidade do usuário", "Propriedades de navegação"],
        "lgpd_risk": "high"
    },
    "amplitude": {
        "name": "Amplitude",
        "patterns": [
            r"cdn\.amplitude\.com",
            r"amplitude\.getInstance",
            r"amplitude\.logEvent"
        ],
        "type": "analytics",
        "data_collected": ["Eventos de produto", "Cohorts", "Jornada do usuário"],
        "lgpd_risk": "high"
    },
    "mixpanel": {
        "name": "Mixpanel",
        "patterns": [
            r"cdn\.mxpnl\.com",
            r"mixpanel\.init",
            r"mixpanel\.track"
        ],
        "type": "analytics",
        "data_collected": ["Eventos de produto", "Perfil do usuário", "Funis"],
        "lgpd_risk": "high"
    },
    "heap": {
        "name": "Heap Analytics",
        "patterns": [
            r"heap-analytics\.com",
            r"heapanalytics\.com",
            r"heap\.load"
        ],
        "type": "analytics",
        "data_collected": ["Auto-tracking de eventos", "Sessões", "Comportamento do usuário"],
        "lgpd_risk": "critical"
    },
    "hubspot": {
        "name": "HubSpot",
        "patterns": [
            r"js\.hs-scripts\.com",
            r"js\.hs-analytics\.net",
            r"hbspt\.forms",
            r"hubspot"
        ],
        "type": "marketing",
        "data_collected": ["Leads", "Formulários", "Tracking de emails", "Chat"],
        "lgpd_risk": "high"
    },
    "pardot": {
        "name": "Salesforce Pardot",
        "patterns": [
            r"pi\.pardot\.com",
            r"pardot\.com/pd\.js",
            r"piAId"
        ],
        "type": "marketing",
        "data_collected": ["Leads B2B", "Scoring", "Email tracking"],
        "lgpd_risk": "high"
    },
    "adobe_analytics": {
        "name": "Adobe Analytics",
        "patterns": [
            r"omtrdc\.net",
            r"sc\.omtrdc\.net",
            r"s\.t\(\)",
            r"AppMeasurement"
        ],
        "type": "analytics",
        "data_collected": ["Pageviews", "Eventos", "eVars", "Props"],
        "lgpd_risk": "high"
    },
    "taboola": {
        "name": "Taboola Pixel",
        "patterns": [
            r"cdn\.taboola\.com",
            r"_tfa\.push",
            r"taboola"
        ],
        "type": "advertising",
        "data_collected": ["Conversões", "Remarketing nativo"],
        "lgpd_risk": "high"
    },
    "outbrain": {
        "name": "Outbrain Pixel",
        "patterns": [
            r"outbrain\.com",
            r"obApi\s*\(",
            r"outbrain"
        ],
        "type": "advertising",
        "data_collected": ["Conversões", "Remarketing nativo"],
        "lgpd_risk": "high"
    },
    "criteo": {
        "name": "Criteo",
        "patterns": [
            r"static\.criteo\.net",
            r"criteo\.com",
            r"criteo_q\.push"
        ],
        "type": "advertising",
        "data_collected": ["Retargeting", "Comportamento de compra", "Produtos visualizados"],
        "lgpd_risk": "critical"
    },
    "quora_pixel": {
        "name": "Quora Pixel",
        "patterns": [
            r"a\.quora\.com",
            r"qp\s*\("
        ],
        "type": "advertising",
        "data_collected": ["Conversões", "Audiências"],
        "lgpd_risk": "high"
    },
    "reddit_pixel": {
        "name": "Reddit Pixel",
        "patterns": [
            r"ads\.reddit\.com",
            r"rdt\s*\("
        ],
        "type": "advertising",
        "data_collected": ["Conversões", "Audiências Reddit"],
        "lgpd_risk": "high"
    },
    "usercentrics": {
        "name": "Usercentrics (CMP)",
        "patterns": [
            r"usercentrics\.eu",
            r"usercentrics\.com",
            r"UC_UI"
        ],
        "type": "consent",
        "data_collected": ["Preferências de consentimento"],
        "lgpd_risk": "none"
    },
    "quantcast": {
        "name": "Quantcast Choice (CMP)",
        "patterns": [
            r"quantcast\.mgr\.consensu\.org",
            r"quantcast\.com",
            r"__cmp"
        ],
        "type": "consent",
        "data_collected": ["Preferências de consentimento"],
        "lgpd_risk": "none"
    },
    # ============ NOVOS TRACKERS & FERRAMENTAS (UPGRADE) ============
    "vwo": {
        "name": "VWO (Visual Website Optimizer)",
        "patterns": [r"dev\.visualwebsiteoptimizer\.com", r"_vwo_uuid"],
        "type": "ab_testing",
        "data_collected": ["Variações de teste", "Gravações de clique", "ID do visitante"],
        "lgpd_risk": "high"
    },
    "optimizely": {
        "name": "Optimizely",
        "patterns": [r"cdn\.optimizely\.com", r"optimizely\s*\("],
        "type": "ab_testing",
        "data_collected": ["Experimentos", "Dados de eventos", "Atributos do usuário"],
        "lgpd_risk": "high"
    },
    "crazyegg": {
        "name": "Crazy Egg",
        "patterns": [r"script\.crazyegg\.com", r"crazyegg"],
        "type": "heatmap",
        "data_collected": ["Mapas de calor", "Sessões"],
        "lgpd_risk": "high"
    },
    "mouseflow": {
        "name": "Mouseflow",
        "patterns": [r"cdn\.mouseflow\.com", r"mouseflow\s*\("],
        "type": "heatmap",
        "data_collected": ["Gravação de tela", "Keystrokes (se não mascarado)", "Movimento do mouse"],
        "lgpd_risk": "critical"
    },
    "activecampaign": {
        "name": "ActiveCampaign",
        "patterns": [r"trackcmp\.net", r"prism\.app"],
        "type": "marketing",
        "data_collected": ["Tracking de email", "Visitas ao site", "Perfil de CRM"],
        "lgpd_risk": "high"
    },
    "mailchimp": {
        "name": "Mailchimp",
        "patterns": [r"chimpstatic\.com", r"mc-validate"],
        "type": "marketing",
        "data_collected": ["Inscrições em newsletter", "Dados de formulário"],
        "lgpd_risk": "medium"
    },
    "outbrain": {
        "name": "Outbrain",
        "patterns": [r"widgets\.outbrain\.com", r"obApi"],
        "type": "advertising",
        "data_collected": ["Interesse de conteúdo", "Remarketing"],
        "lgpd_risk": "high"
    },
    "taboola": {
        "name": "Taboola",
        "patterns": [r"cdn\.taboola\.com", r"taboola"],
        "type": "advertising",
        "data_collected": ["Preferências de leitura", "Remarketing nético"],
        "lgpd_risk": "high"
    },
    "braze": {
        "name": "Braze",
        "patterns": [r"sdk\.braze\.com", r"braze\s*\("],
        "type": "marketing",
        "data_collected": ["Notificações push", "Atributos de usuário", "Eventos de app"],
        "lgpd_risk": "high"
    },
    "klaviyo": {
        "name": "Klaviyo",
        "patterns": [r"static\.klaviyo\.com", r"klaviyo\s*\("],
        "type": "marketing",
        "data_collected": ["Dados de e-commerce", "Tracking de comportamento de compra"],
        "lgpd_risk": "high"
    },
    "drift": {
        "name": "Drift",
        "patterns": [r"js\.driftt\.com", r"driftt"],
        "type": "support",
        "data_collected": ["Conversas de chat", "Identificação de empresa (IP)"],
        "lgpd_risk": "high"
    },
    "crisp": {
        "name": "Crisp",
        "patterns": [r"client\.crisp\.chat"],
        "type": "support",
        "data_collected": ["Mensagens", "Localização do visitante"],
        "lgpd_risk": "medium"
    },
    "luckyorange": {
        "name": "Lucky Orange",
        "patterns": [r"w\.luckyorange\.com", r"luckyorange"],
        "type": "heatmap",
        "data_collected": ["Gravações", "Pó-formulário"],
        "lgpd_risk": "high"
    },
    "fullstory": {
        "name": "FullStory",
        "patterns": [r"fullstory\.com/s/fs\.js"],
        "type": "heatmap",
        "data_collected": ["Gravação de alta fidelidade", "Console logs do usuário"],
        "lgpd_risk": "critical"
    }
}

# Generic Privacy Risks mapping
PRIVACY_RISKS = {
    "consent": {
        "article": "Consent Requirement",
        "description": "Tracking before explicit consent",
        "fine_range": (5000, 20000)
    },
    "data_sharing": {
        "article": "Data Sharing & Advertising",
        "description": "Third-party sharing without proper disclosure",
        "fine_range": (10000, 50000)
    },
    "sensitive_data": {
        "article": "High-Risk Tracking",
        "description": "Use of intrusive/critical tracking tools",
        "fine_range": (15000, 50000)
    },
    "data_protection": {
        "article": "Data Security",
        "description": "Use of obsolete or insecure tracking methods",
        "fine_range": (5000, 30000)
    },
    "transparency": {
        "article": "Transparency",
        "description": "Lack of mechanisms for user choice (CMP)",
        "fine_range": (2000, 10000)
    }
}


def load_db() -> List[Dict]:
    if not os.path.isfile(DB_PATH):
        with open(DB_PATH, "w") as f:
            json.dump([], f)
    with open(DB_PATH, "r", encoding="utf-8") as f:
        return json.load(f)


def save_db(data: List[Dict]) -> None:
    with open(DB_PATH, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)


def detect_tags(html: str) -> List[Dict[str, Any]]:
    """Detecta tags de tracking no HTML com posição real"""
    detected_tags = []
    seen_tags = set()
    
    for tag_id, tag_info in TAG_PATTERNS.items():
        for pattern in tag_info["patterns"]:
            matches = list(re.finditer(pattern, html, re.IGNORECASE))
            for match in matches:
                # Calcular posição no documento
                position = match.start()
                line_number = html[:position].count('\n') + 1
                
                # Extrair contexto
                start = max(0, position - 100)
                end = min(len(html), position + 200)
                context = html[start:end]
                
                # Extrair ID e detectar duplicatas
                tag_match_id = match.group() if len(match.group()) < 30 else None
                
                # Check for specific GTM or GA4 IDs to track exact duplicates
                exact_id = ""
                if tag_id == "gtm":
                    gtm_match = re.search(r"GTM-[A-Z0-9]+", match.group(), re.IGNORECASE)
                    if gtm_match: exact_id = gtm_match.group(0)
                elif tag_id == "ga4":
                    ga4_match = re.search(r"G-[A-Z0-9]+", match.group(), re.IGNORECASE)
                    if ga4_match: exact_id = ga4_match.group(0)
                    
                tag_key = f"{tag_id}_{exact_id if exact_id else match.group()[:20]}"
                
                is_duplicate = False
                if tag_key in seen_tags:
                    is_duplicate = True
                else:
                    seen_tags.add(tag_key)
                
                # Detectar ID específico se existir
                tag_match_id = match.group() if len(match.group()) < 30 else None
                
                detected_tags.append({
                    "id": tag_id,
                    "name": tag_info["name"],
                    "type": tag_info["type"],
                    "position": position,
                    "lineNumber": line_number,
                    "matchedPattern": match.group()[:50],
                    "tagId": exact_id if exact_id else tag_match_id,
                    "dataCollected": tag_info["data_collected"],
                    "lgpdRisk": tag_info["lgpd_risk"],
                    "context": context[:150],
                    "isDuplicate": is_duplicate,
                    "isBeforeConsent": False  # Será calculado depois
                })
    
    # Ordenar por posição
    detected_tags.sort(key=lambda x: x["position"])
    
    # Marcar tags antes do consentimento
    consent_position = None
    for tag in detected_tags:
        if tag["type"] == "consent":
            consent_position = tag["position"]
            break
    
    if consent_position:
        for tag in detected_tags:
            if tag["position"] < consent_position and tag["type"] != "consent":
                tag["isBeforeConsent"] = True
    else:
        # Se não há CMP, todas as tags de tracking violam LGPD
        for tag in detected_tags:
            if tag["type"] in ["analytics", "advertising", "heatmap"]:
                tag["isBeforeConsent"] = True
    
    return detected_tags


def analyze_privacy_compliance(tags: List[Dict], jurisdiction: Dict) -> Dict[str, Any]:
    """Calcula riscos de privacidade com base no comportamento observável e jurisdição."""
    violations = []
    total_min_fine = 0
    total_max_fine = 0
    
    has_consent_tool = any(t["type"] == "consent" for t in tags)
    has_ua = any(t["id"] == "ua" for t in tags)
    
    consent_risks = []
    disclosure_gaps = []
    
    for tag in tags:
        if tag["isBeforeConsent"] and tag["type"] != "consent":
            violation_type = "consent"
            if tag["type"] == "advertising":
                violation_type = "data_sharing"
            elif tag["lgpdRisk"] == "critical":
                violation_type = "sensitive_data"
            
            risk_info = PRIVACY_RISKS[violation_type]
            if violation_type == "consent":
                consent_risks.append(f"{tag['name']} firing before consent")
            
            violations.append({
                "tag": tag["name"],
                "tagId": tag.get("tagId"),
                "violation": f"Carregado antes do consentimento (linha {tag['lineNumber']})",
                "article": risk_info["article"],
                "description": risk_info["description"],
                "severity": tag["lgpdRisk"],
                "dataCollected": tag["dataCollected"],
                "estimatedFine": "Indicativo de não conformidade"
            })
            total_min_fine += risk_info["fine_range"][0]
            total_max_fine += risk_info["fine_range"][1]
    
    # Risco adicional se não tem CMP
    if not has_consent_tool and any(t["type"] in ["analytics", "advertising"] for t in tags):
        risk_info = PRIVACY_RISKS["transparency"]
        disclosure_gaps.append("Missing Consent Management Platform (CMP)")
        violations.append({
            "tag": "Sistema",
            "tagId": None,
            "violation": "Nenhuma ferramenta de consentimento (CMP) detectada",
            "article": risk_info["article"],
            "description": risk_info["description"],
            "severity": "critical",
            "dataCollected": [],
            "estimatedFine": "Alto risco de conformidade"
        })
        total_min_fine += risk_info["fine_range"][0]
        total_max_fine += risk_info["fine_range"][1]
    
    # Risco adicional se usa Universal Analytics (obsoleto)
    if has_ua:
        disclosure_gaps.append("Using obsolete Universal Analytics")
        violations.append({
            "tag": "Universal Analytics",
            "tagId": None,
            "violation": "Uso de tecnologia obsoleta (UA descontinuado em julho/2023)",
            "article": "Data Security Risk",
            "description": "Falha em manter medidas de segurança atualizadas",
            "severity": "critical",
            "dataCollected": ["Dados históricos sem controle"],
            "estimatedFine": "Vulnerabilidade técnica detectada"
        })
        total_min_fine += 10000
        total_max_fine += 30000
        
    # Verificação de Duplicatas e Erros
    duplicates_count = 0
    for t in tags:
        if t.get("isDuplicate"):
            duplicates_count += 1
            disclosure_gaps.append(f"Duplicate Tag: {t.get('name')} ({t.get('tagId') or 'N/A'})")
            violations.append({
                "tag": t.get("name"),
                "tagId": t.get("tagId"),
                "violation": "Implementação Duplicada Detectada",
                "article": "Data Minimization / Performance",
                "description": f"A tag foi inserida múltiplas vezes na página, causando disparos duplos e inconsistência de dados.",
                "severity": "high",
                "dataCollected": t.get("dataCollected", []),
                "estimatedFine": "Risco analítico alto"
            })
            total_max_fine += 5000
    
    estimated_risk_exposure = f"Exposição Potencial (apenas ref.): USD ${int(total_min_fine/5):,} - USD ${int(total_max_fine/5):,}" if violations else "Baixo Risco"
    
    return {
        "jurisdiction": jurisdiction["region"],
        "framework": jurisdiction["law_full"],
        "confidenceLevel": "Avaliação baseada no comportamento público do site (não constitui conselho legal)",
        "consentRisks": list(set(consent_risks)),
        "disclosureGaps": list(set(disclosure_gaps)),
        "violations": violations,
        "totalViolations": len(violations),
        "estimatedRiskExposure": estimated_risk_exposure,
        "hasConsentTool": has_consent_tool,
        "hasUniversalAnalytics": has_ua,
        "complianceScore": max(0, 100 - (len(violations) * 12)) 
    }


def analyze_datalayer_structure(html: str) -> List[Dict[str, Any]]:
    """Detecta a presença e chaves do dataLayer para identificar vazamentos de PII"""
    datalayer_findings = []
    
    # Padrão para dataLayer.push
    push_pattern = r"dataLayer\.push\s*\(\s*(\{.*?\})\s*\)"
    matches = re.finditer(push_pattern, html, re.DOTALL)
    
    pii_keywords = ['email', 'nome', 'cpf', 'telefone', 'phone', 'birth', 'nascimento', 'password', 'senha', 'user_id', 'uid']
    
    for match in matches:
        content = match.group(1)
        found_pii = []
        for key in pii_keywords:
            if key in content.lower():
                found_pii.append(key)
        
        if found_pii:
            datalayer_findings.append({
                "type": "pii_exposure",
                "keys": found_pii,
                "snippet": content[:100] + "...",
                "severity": "critical" if any(k in found_pii for k in ['cpf', 'password', 'email']) else "high"
            })
            
    return datalayer_findings


def check_script_security(html: str) -> List[Dict[str, Any]]:
    """Verifica segurança de scripts externos (SRI, Protocolo)"""
    soup = BeautifulSoup(html, 'html.parser')
    security_issues = []
    
    for script in soup.find_all('script', src=True):
        src = script['src']
        
        # Check for SRI
        if 'integrity' not in script.attrs and not src.startswith(('/', 'http://localhost', 'https://localhost')):
            security_issues.append({
                "type": "missing_sri",
                "src": src[:50] + "...",
                "severity": "low",
                "description": "Script externo sem integridade (SRI) detectada"
            })
            
        # Check for HTTP (insecure)
        if src.startswith('http://'):
            security_issues.append({
                "type": "insecure_protocol",
                "src": src,
                "severity": "high",
                "description": "Script carregado via HTTP inseguro"
            })
            
    return security_issues


def analyze_loading_order(tags: List[Dict]) -> List[Dict[str, Any]]:
    """Analisa ordem de carregamento e recomendações"""
    order_analysis = []
    
    # Ordem ideal
    ideal_order = ["consent", "tag_manager", "analytics", "advertising", "heatmap", "support"]
    
    for i, tag in enumerate(tags):
        issues = []
        recommendations = []
        
        # Verificar se CMP está primeiro
        if tag["type"] == "consent" and i > 0:
            issues.append("CMP deve ser a primeira tag a carregar")
            recommendations.append("Mover para o início do <head>")
        
        # Verificar se analytics/advertising vem antes do consent
        if tag["type"] in ["analytics", "advertising", "heatmap"]:
            consent_index = next((j for j, t in enumerate(tags) if t["type"] == "consent"), -1)
            if consent_index == -1 or i < consent_index:
                issues.append("Tag de tracking antes do consentimento")
                recommendations.append("Implementar via GTM com trigger de consentimento")
        
        order_analysis.append({
            "position": i + 1,
            "tag": tag["name"],
            "type": tag["type"],
            "lineNumber": tag["lineNumber"],
            "issues": issues,
            "recommendations": recommendations,
            "severity": "critical" if issues else "ok"
        })
    
    return order_analysis


def detect_consent_mode_v2(html: str) -> Dict[str, Any]:
    """Detecta configuração do Google Consent Mode v2"""
    consent_mode = {
        "detected": False,
        "version": None,
        "defaultState": {},
        "hasDefaultDenied": False,
        "signals": [],
        "issues": [],
        "recommendations": []
    }
    
    # Detectar gtag consent default
    default_pattern = r"gtag\s*\(\s*['\"]consent['\"]\s*,\s*['\"]default['\"]\s*,\s*(\{[^}]+\})"
    default_match = re.search(default_pattern, html, re.IGNORECASE | re.DOTALL)
    
    if default_match:
        consent_mode["detected"] = True
        config_str = default_match.group(1)
        
        # Detectar sinais v2
        v2_signals = ["ad_user_data", "ad_personalization"]
        for signal in v2_signals:
            if signal in config_str:
                consent_mode["signals"].append(signal)
                consent_mode["version"] = "v2"
        
        # Sinais básicos
        basic_signals = ["ad_storage", "analytics_storage"]
        for signal in basic_signals:
            if signal in config_str:
                consent_mode["signals"].append(signal)
                if consent_mode["version"] != "v2":
                    consent_mode["version"] = "v1"
        
        # Verificar estado default
        for signal in ["ad_storage", "analytics_storage", "ad_user_data", "ad_personalization"]:
            denied_pattern = rf"['\"]?{signal}['\"]?\s*:\s*['\"]denied['\"]"
            granted_pattern = rf"['\"]?{signal}['\"]?\s*:\s*['\"]granted['\"]"
            if re.search(denied_pattern, config_str):
                consent_mode["defaultState"][signal] = "denied"
                consent_mode["hasDefaultDenied"] = True
            elif re.search(granted_pattern, config_str):
                consent_mode["defaultState"][signal] = "granted"
        
        # Validações
        if not consent_mode["hasDefaultDenied"]:
            consent_mode["issues"].append("Consent Mode não está com default 'denied' - possível violação LGPD")
            consent_mode["recommendations"].append("Configurar todos os sinais como 'denied' por padrão")
        
        if consent_mode["version"] != "v2":
            consent_mode["issues"].append("Consent Mode v1 detectado - Google exige v2 desde março/2024")
            consent_mode["recommendations"].append("Adicionar ad_user_data e ad_personalization para compliance v2")
        
        if "wait_for_update" not in config_str:
            consent_mode["recommendations"].append("Adicionar 'wait_for_update': 500 para aguardar resposta do CMP")
    
    else:
        # Verificar se há GA4 ou Google Ads sem Consent Mode
        if re.search(r"gtag\s*\(\s*['\"]config['\"]", html) or re.search(r"googletagmanager\.com/gtag", html):
            consent_mode["issues"].append("Google Analytics/Ads detectado SEM Consent Mode configurado")
            consent_mode["recommendations"].append("Implementar Consent Mode v2 antes do gtag config")
    
    return consent_mode



def detect_jurisdiction(html: str, url: str) -> Dict[str, Any]:
    """
    Detects likely jurisdiction based on site content, language, and TLD.
    Does NOT guarantee legal compliance. Returns indicative assessment only.
    """
    parsed = urlparse(url)
    tld = parsed.netloc.split('.')[-1].lower()
    domain = parsed.netloc.lower()

    # TLD-based hints
    if tld in ('br', 'com.br'):
        return {"region": "Brazil", "law": "LGPD", "law_full": "Lei Geral de Proteção de Dados (LGPD) — Brazil"}
    if tld in ('uk', 'co.uk'):
        return {"region": "United Kingdom", "law": "UK GDPR", "law_full": "UK General Data Protection Regulation"}
    if tld in ('eu', 'de', 'fr', 'it', 'es', 'pt', 'nl', 'be', 'at', 'pl', 'se', 'fi', 'dk', 'ie'):
        return {"region": "EU / EEA", "law": "GDPR", "law_full": "General Data Protection Regulation (GDPR) — EU/EEA"}
    if tld in ('us', 'gov', 'edu') or domain.endswith('.ca.gov'):
        return {"region": "United States", "law": "CCPA/CPRA", "law_full": "California Consumer Privacy Act / US State Privacy Laws"}

    # Content-based language detection
    pt_patterns = [r'\bportuguês\b', r'\bpolítica de privacidade\b', r'\bcookies\b.*\baceitar\b', r'\bLGPD\b']
    de_en_eu_patterns = [r'\bDatenschutz\b', r'\bCookie-Einstellungen\b', r'\bRGPD\b', r'\bGDPR\b']

    for p in pt_patterns:
        if re.search(p, html, re.IGNORECASE):
            return {"region": "Brazil", "law": "LGPD", "law_full": "Lei Geral de Proteção de Dados (LGPD) — Brazil"}

    for p in de_en_eu_patterns:
        if re.search(p, html, re.IGNORECASE):
            return {"region": "EU / EEA", "law": "GDPR", "law_full": "General Data Protection Regulation (GDPR) — EU/EEA"}

    # Default: general risk assessment
    return {"region": "Global / Other", "law": "General Privacy", "law_full": "General Privacy Risk Assessment"}


def generate_ai_recommendations(tags: List[Dict], lgpd_analysis: Dict, consent_mode: Dict,
                                 datalayer: List, security: List, jurisdiction: Dict, html: str) -> List[Dict]:
    """Generates rule-based actionable recommendations per finding."""
    recs = []
    law = jurisdiction.get("law", "Privacy Law")

    has_ga4 = any(t["id"] == "ga4" for t in tags)
    has_gtm = any(t["id"] == "gtm" for t in tags)
    has_meta = any(t["id"] == "meta_pixel" for t in tags)
    has_consent = lgpd_analysis.get("hasConsentTool", False)
    has_ua = lgpd_analysis.get("hasUniversalAnalytics", False)
    consent_v2 = consent_mode.get("version") == "v2"
    has_default_denied = consent_mode.get("hasDefaultDenied", False)

    if not has_consent:
        recs.append({
            "severity": "critical",
            "category": "consent",
            "title": "No consent management platform (CMP) detected",
            "detail": f"No cookie banner or consent manager was found on the scanned page. Under {law}, tracking scripts must not fire before explicit user consent.",
            "fix": "Implement a CMP (e.g., Cookiebot, OneTrust, LGPD-compliant banner) as the first element loading in <head>."
        })

    if has_ga4 and not consent_v2:
        recs.append({
            "severity": "high",
            "category": "consent_mode",
            "title": "Google Analytics 4 detected without Consent Mode v2",
            "detail": "GA4 was found but Google Consent Mode v2 is not configured. Google required CMv2 compliance from March 2024 for EU users.",
            "fix": "Add gtag('consent','default',{...}) with all signals set to 'denied' before the GA4 config call."
        })

    if has_meta and not has_consent:
        recs.append({
            "severity": "critical",
            "category": "tracking_before_consent",
            "title": "Meta Pixel loading before consent",
            "detail": f"Meta Pixel (Facebook) was detected loading before any consent signal. This likely violates {law}.",
            "fix": "Move Meta Pixel to fire via GTM using a consent trigger, or block it until the user accepts cookies."
        })

    if not has_gtm and (has_ga4 or has_meta):
        recs.append({
            "severity": "medium",
            "category": "implementation",
            "title": "Tracking tags found without Google Tag Manager",
            "detail": "GA4/Meta Pixel were detected but no GTM container was found. Managing tags directly in code makes consent orchestration very difficult.",
            "fix": "Implement Google Tag Manager and migrate all tracking scripts into it for centralized consent management."
        })

    if has_ua:
        recs.append({
            "severity": "critical",
            "category": "obsolete_technology",
            "title": "Universal Analytics (UA) detected — deprecated since July 2023",
            "detail": "UA was discontinued by Google in July 2023. Continued use may result in data collection using unmaintained code.",
            "fix": "Migrate all UA tracking to Google Analytics 4 (GA4) immediately."
        })

    if consent_mode.get("detected") and not has_default_denied:
        recs.append({
            "severity": "high",
            "category": "consent_mode",
            "title": "Consent Mode detected but not set to 'denied' by default",
            "detail": "Google Consent Mode was found but the default state is not 'denied'. This means tracking may fire before user consent.",
            "fix": "Set all consent signals (ad_storage, analytics_storage, ad_user_data, ad_personalization) to 'denied' by default."
        })

    for dl in datalayer:
        recs.append({
            "severity": "critical",
            "category": "pii_exposure",
            "title": f"PII data detected in dataLayer: {', '.join(dl.get('keys', []))}",
            "detail": "Personal Identifiable Information (PII) was found being pushed to the dataLayer, which may be readable by all loaded tracking scripts.",
            "fix": "Remove PII from dataLayer pushes. Use hashed user IDs instead of raw email/CPF/phone."
        })

    high_security = [s for s in security if s.get("severity") == "high"]
    for s in high_security[:2]:
        recs.append({
            "severity": "high",
            "category": "security",
            "title": "Script loaded via insecure HTTP",
            "detail": f"Script {s.get('src', 'unknown')[:60]} loads over HTTP, making it vulnerable to MITM injection.",
            "fix": "Update all external script references to use HTTPS."
        })

    # Check for privacy policy
    if not re.search(r'privacy.?policy|política.?de.?privacidade|datenschutz', html if isinstance(html, str) else '', re.IGNORECASE):
        recs.append({
            "severity": "medium",
            "category": "legal_pages",
            "title": "Privacy policy page not detected",
            "detail": f"No link to a privacy policy was found in the scanned page. A privacy policy is required by {law}.",
            "fix": "Add a clearly visible link to your privacy policy in the page footer."
        })

    return recs


def populate_scores(lgpd_analysis: Dict, consent_mode: Dict,
                    datalayer: List, security: List, tags: List) -> Dict[str, Any]:
    """Generates 6 risk scores."""
    base_score = lgpd_analysis.get("complianceScore", 50)
    violations = lgpd_analysis.get("totalViolations", 0)

    # Tracking Quality Score
    has_gtm = any(t["id"] == "gtm" for t in tags)
    has_ga4 = any(t["id"] == "ga4" for t in tags)
    has_ua = lgpd_analysis.get("hasUniversalAnalytics", False)
    dup_count = sum(1 for t in tags if tags.count(t) > 1)

    tracking_quality = 100
    if not has_gtm:
        tracking_quality -= 25
    if has_ua:
        tracking_quality -= 30
    if not has_ga4:
        tracking_quality -= 15
    tracking_quality = max(0, tracking_quality - (dup_count * 5))

    # Privacy Risk Score (higher = more risk)
    privacy_risk = min(100, violations * 15 + (0 if lgpd_analysis.get("hasConsentTool") else 30))

    # Consent Integrity Score
    consent_score = 100
    if not lgpd_analysis.get("hasConsentTool"):
        consent_score -= 60
    if not consent_mode.get("detected"):
        consent_score -= 20
    if not consent_mode.get("hasDefaultDenied"):
        consent_score -= 15
    if consent_mode.get("version") != "v2":
        consent_score -= 10
    consent_score = max(0, consent_score)

    # DataLayer Quality Score
    datalayer_score = 100 - (len(datalayer) * 25)
    datalayer_score = max(0, datalayer_score)

    # Tag Duplication Risk
    tag_dup_risk = min(100, dup_count * 20)

    # Estimated Business Impact
    if privacy_risk >= 70:
        impact = "high"
    elif privacy_risk >= 35:
        impact = "medium"
    else:
        impact = "low"

    return {
        "auditScore": base_score,
        "trackingQualityScore": tracking_quality,
        "privacyRiskScore": privacy_risk,
        "consentIntegrityScore": consent_score,
        "datalayerQualityScore": datalayer_score,
        "tagDuplicationRisk": tag_dup_risk,
        "estimatedBusinessImpact": impact,
    }


def generate_audit_report(scan_results: Dict[str, Any]) -> Dict[str, Any]:
    """Gera relatório completo de auditoria com jurisdição, scores e recomendações."""
    html = scan_results.get("html", "")
    url = scan_results.get("url", "")
    
    tags = detect_tags(html)
    
    # Enrich with network requests from Playwright
    network_requests = scan_results.get("network_requests", [])
    for t in tags:
        # Check if this tag actually fired a network request
        fired = any(t.get("matchedPattern") in req["url"] or t["id"] in req["url"] for req in network_requests)
        t["networkFired"] = fired

    jurisdiction = detect_jurisdiction(html, url)
    privacy_analysis = analyze_privacy_compliance(tags, jurisdiction)
    order_analysis = analyze_loading_order(tags)
    consent_mode_analysis = detect_consent_mode_v2(html)
    
    # Merging static dataLayer analysis with real extracted dataLayer from Playwright
    datalayer_analysis = analyze_datalayer_structure(html)
    real_dl = scan_results.get("dataLayer", [])
    if real_dl:
        # Add a node representing real extracted JSON
        datalayer_analysis.append({
            "type": "runtime_dataLayer",
            "keys": ["Extracted live object count: " + str(len(real_dl))],
            "snippet": json.dumps(real_dl)[:200] + "...",
            "severity": "info"
        })
        
    security_analysis = check_script_security(html)
    jurisdiction = detect_jurisdiction(html, url)
    privacy_analysis = analyze_privacy_compliance(tags, jurisdiction)
    order_analysis = analyze_loading_order(tags)
    consent_mode_analysis = detect_consent_mode_v2(html)
    datalayer_analysis = analyze_datalayer_structure(html)
    security_analysis = check_script_security(html)
    recommendations = generate_ai_recommendations(
        tags, privacy_analysis, consent_mode_analysis, datalayer_analysis, security_analysis, jurisdiction, html
    )
    scores = populate_scores(privacy_analysis, consent_mode_analysis, datalayer_analysis, security_analysis, tags)

    # Contar por tipo
    type_counts = {}
    for tag in tags:
        type_counts[tag["type"]] = type_counts.get(tag["type"], 0) + 1
    
    # Detectar duplicatas
    duplicates = []
    tag_names = [t["name"] for t in tags]
    for name in set(tag_names):
        count = tag_names.count(name)
        if count > 1:
            duplicates.append({
                "tag": name,
                "count": count,
                "severity": "warning",
                "recommendation": f"Consolidar {count} instâncias de {name}"
            })
    
    # Ajustar score com base no Consent Mode e Segurança
    final_score = privacy_analysis["complianceScore"]
    
    # Bonus/Penalidade Consent Mode
    if consent_mode_analysis["detected"] and consent_mode_analysis["hasDefaultDenied"]:
        final_score = min(100, final_score + 10)
    elif consent_mode_analysis["issues"]:
        final_score = max(0, final_score - 10)
        
    # Penalidade por PII exposto no DataLayer
    if datalayer_analysis:
        final_score = max(0, final_score - (len(datalayer_analysis) * 15))
        
    # Penalidade por Segurança Insegura
    high_security_issues = [s for s in security_analysis if s["severity"] == "high"]
    if high_security_issues:
        final_score = max(0, final_score - 20)
    
    return {
        "url": url,
        "timestamp": datetime.now().isoformat(),
        "tags": tags,
        "tagCount": len(tags),
        "typeCounts": type_counts,
        "duplicates": duplicates,
        "loadingOrder": order_analysis,
        "privacy": privacy_analysis, # Kept for API compatibility with frontend NewAuditResults
        "consentModeAnalysis": consent_mode_analysis,
        "datalayerAnalysis": datalayer_analysis,
        "securityAnalysis": security_analysis,
        "jurisdiction": jurisdiction,
        "scores": scores,
        "recommendations": recommendations,
        "recommendationSummary": {
            "critical": sum(1 for r in recommendations if r.get("severity") == "critical"),
            "high": sum(1 for r in recommendations if r.get("severity") == "high"),
            "medium": sum(1 for r in recommendations if r.get("severity") == "medium"),
            "low": sum(1 for r in recommendations if r.get("severity") == "low"),
        },
        "score": final_score,
        "summary": {
            "totalTags": len(tags),
            "consentDetected": privacy_analysis["hasConsentTool"],
            "consentModeV2": consent_mode_analysis["version"] == "v2",
            "hasUniversalAnalytics": privacy_analysis["hasUniversalAnalytics"],
            "violationsCount": privacy_analysis["totalViolations"],
            "estimatedFine": privacy_analysis.get("estimatedRiskExposure", ""),
            "hasDatalayerPII": len(datalayer_analysis) > 0,
            "securityIssuesCount": len(security_analysis)
        },
        "disclaimer": "This platform provides an indicative technical and privacy risk assessment based on publicly observable site behavior. It does not replace formal legal advice."
    }



# ============================================================
# SMART URL FETCHER — handles 403 / Cloudflare / CDN blocks
# ============================================================

class BlockedByRobots(Exception):
    """Raised when a site actively blocks our scraper."""
    pass

_BROWSER_AGENTS = [
    # Chrome 124 Windows
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
    # Chrome 124 macOS
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
    # Firefox 125 Windows
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:125.0) Gecko/20100101 Firefox/125.0",
]

def _build_headers(ua: str, url: str) -> Dict[str, str]:
    """Build browser-like request headers to pass Cloudflare / CDN checks."""
    parsed = urlparse(url)
    origin = f"{parsed.scheme}://{parsed.netloc}"
    return {
        "User-Agent": ua,
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8",
        "Accept-Language": "pt-BR,pt;q=0.9,en-US;q=0.8,en;q=0.7",
        "Accept-Encoding": "gzip, deflate, br",
        "Referer": origin + "/",
        "Origin": origin,
        "Connection": "keep-alive",
        "Upgrade-Insecure-Requests": "1",
        "Sec-Fetch-Dest": "document",
        "Sec-Fetch-Mode": "navigate",
        "Sec-Fetch-Site": "same-origin",
        "Sec-Fetch-User": "?1",
        "Cache-Control": "max-age=0",
        "DNT": "1",
    }

def smart_fetch(url: str, timeout: int = 20) -> str:
    """
    Fetches a URL with browser-like headers and retry on 403/429/503.
    Returns HTML string. Raises BlockedByRobots if all attempts fail.
    """
    session = requests.Session()
    last_status = 0
    last_error = None

    for idx, ua in enumerate(_BROWSER_AGENTS):
        try:
            resp = session.get(
                url,
                headers=_build_headers(ua, url),
                timeout=timeout,
                allow_redirects=True,
                verify=False
            )
            last_status = resp.status_code

            if resp.status_code == 200:
                resp.encoding = resp.apparent_encoding or "utf-8"
                return resp.text

            if resp.status_code in (401, 403, 429, 503):
                print(f" Attempt {idx+1}: got {resp.status_code} for {url}, retrying...")
                continue  # try next UA

            # Other 4xx/5xx — raise immediately
            resp.raise_for_status()

        except requests.RequestException as e:
            last_error = e
            print(f" Attempt {idx+1} network error for {url}: {e}")
            continue

    # All attempts failed
    if last_status in (401, 403):
        raise BlockedByRobots(
            f"O site {url} bloqueia solicitações automáticas (HTTP {last_status}). "
            "Tente usar a aba Upload para analisar o GTM Container JSON exportado, "
            "ou use a aba View-Source depois de abrir a página no navegador."
        )

    if last_error:
        raise last_error

    raise BlockedByRobots(f"Não foi possível acessar {url} após {len(_BROWSER_AGENTS)} tentativas.")


# ============================================================
# ROUTES
# ============================================================


@app.get("/")
def root():
    index_path = os.path.join(STATIC_DIR, "index.html")
    if os.path.exists(index_path):
        return FileResponse(index_path)
    # Fallback se não existir dist
    return JSONResponse({
        "status": "ok",
        "message": "DataTrust Audit API v2.0",
        "docs": "/docs"
    })


@app.get("/health")
def health():
    return JSONResponse({
        "status": "healthy",
        "version": "2.0.0",
        "timestamp": datetime.now().isoformat()
    })


def plan_strip_report(report: dict, plan: str) -> None:
    """Strip report payload server-side based on plan entitlement.
    
    Mutates `report` in-place.
    
    free  → 30% tags (no lineNumber/context/matchedPattern), violation count only, redacted fine
    pro   → 70% tags (violation details OK), still no lineNumber / view-source data
    full  → no stripping
    """
    if plan == "full":
        return  # Premium users get everything

    tags = report.get("tags", [])
    visible_pct = 0.30 if plan == "free" else 0.70
    visible_count = max(1, int(len(tags) * visible_pct))

    # Truncate tag list
    stripped_tags = tags[:visible_count]

    # Strip sensitive fields per plan
    for tag in stripped_tags:
        # Line number and code position — Premium only
        tag.pop("lineNumber", None)
        tag.pop("context", None)

        if plan == "free":
            # Free also strips the matched pattern (exact code snippet)
            tag.pop("matchedPattern", None)
            tag.pop("dataCollected", None)
            tag["lgpdRisk"] = tag.get("lgpdRisk", "unknown")  # keep risk level for visibility

    report["tags"] = stripped_tags

    if plan == "free":
        # Replace violation details with count-only metadata
        lgpd = report.get("privacy", {})
        violation_count = len(lgpd.get("violations", []))
        lgpd["violations"] = []          # empty — details locked
        lgpd["totalViolations"] = violation_count
        lgpd["estimatedRiskExposure"] = "Risco Moderado"  # redacted
        lgpd["_lockedForPlan"] = "pro"
        report["privacy"] = lgpd

        # Redact summary fine
        summary = report.get("summary", {})
        summary["estimatedFine"] = "Desbloqueie no plano Pro"
        report["summary"] = summary

    # Both free and pro: remove any view-source / deep analysis extras
    report.pop("viewSourceInfo", None)


# Dependency for scanner
from audit_engine.scanner import PlaywrightAuditScanner
scanner = PlaywrightAuditScanner()

@app.post("/analyze")
async def analyze_url(request: Request):
    try:
        data = await request.json()
        target_url = data.get("url", "").strip()
        plan = data.get("plan", "free")
        
        if not target_url:
            return JSONResponse({"status": "error", "msg": "URL não fornecida"}, status_code=400)
            
        if not target_url.startswith(("http://", "https://")):
            target_url = "https://" + target_url
            
        scan_results = await scanner.scan(target_url)
        if scan_results.get("error") and not scan_results.get("html"):
            return JSONResponse({"status": "error", "msg": f"Fala ao carregar página: {scan_results['error']}"}, status_code=500)
            
        report = generate_audit_report(scan_results)
        
        # Add raw metrics for premium users if needed
        report["metrics"] = {
            "consoleWarnings": len(scan_results.get("console_messages", [])),
            "networkRequests": len(scan_results.get("network_requests", []))
        }
        
        # Strip details if plan is free
        plan_strip_report(report, plan)
        
        return JSONResponse(report)

        
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JSONResponse({"status": "error", "msg": str(e)}, status_code=500)
        

@app.post("/deep-analyze")
async def deep_analyze(request: Request):
    """Realiza auditoria em múltiplas páginas do mesmo domínio (Real Deep Scan)"""
    try:
        data = await request.json()
        base_url = data.get("url", "").strip()
        
        if not base_url:
            return JSONResponse({"status": "error", "msg": "URL não fornecida"}, status_code=400)
            
        if not base_url.startswith(("http://", "https://")):
            base_url = "https://" + base_url
            
        # 1. Auditar página inicial
        scan_results = await scanner.scan(base_url)
        base_report = generate_audit_report(scan_results)
        soup = BeautifulSoup(scan_results.get("html", ""), 'html.parser')
        domain = urlparse(base_url).netloc
        links = []
        for a in soup.find_all('a', href=True):
            href = a['href']
            # Se for link relativo ou no mesmo domínio
            if href.startswith('/') or domain in href:
                full_url = urljoin(base_url, href)
                # Evitar âncoras e a própria base_url
                if '#' not in full_url and full_url != base_url and full_url != base_url + '/':
                    links.append(full_url)
        
        # Pegar as 3 primeiras links únicos e interessantes
        target_links = list(set(links))[:3]
        
        sub_reports = [base_report]
        for link in target_links:
            try:
                sub_res = requests.get(link, timeout=10, headers=headers)
                if sub_res.status_code == 200:
                    sub_reports.append(generate_audit_report(sub_res.text, link))
            except:
                continue
                
        # 3. Consolidar resultados
        total_violations = sum(r["lgpdAnalysis"]["totalViolations"] for r in sub_reports)
        avg_score = sum(r["score"] for r in sub_reports) / len(sub_reports)
        all_tags = []
        for r in sub_reports:
            all_tags.extend(r["tags"])
            
        # Remover duplicatas de tags baseadas em nome
        unique_tags = {t["name"]: t for t in all_tags}.values()
        
        deep_report = {
            "url": base_url,
            "timestamp": datetime.now().isoformat(),
            "pagesAudited": len(sub_reports),
            "urls": [r["url"] for r in sub_reports],
            "totalViolations": total_violations,
            "averageScore": round(avg_score, 1),
            "uniqueTagsFound": len(unique_tags),
            "consolidatedTags": list(unique_tags),
            "pageReports": sub_reports
        }
        
        return JSONResponse({
            "status": "ok",
            "result": deep_report
        })
        
    except Exception as e:
        return JSONResponse({"status": "error", "msg": str(e)}, status_code=500)


@app.get("/history")
def history():
    """Retorna histórico de auditorias"""
    history_data = load_db()
    return JSONResponse({
        "history": history_data[::-1],  # Mais recentes primeiro
        "total": len(history_data)
    })


@app.delete("/history")
def clear_history():
    """Limpa histórico de auditorias"""
    save_db([])
    return JSONResponse({
        "status": "ok",
        "message": "Histórico limpo"
    })


@app.get("/lgpd-info")
def lgpd_info():
    """Retorna informações sobre artigos LGPD"""
    return JSONResponse({
        "articles": LGPD_ARTICLES,
        "patterns": {k: {
            "name": v["name"],
            "type": v["type"],
            "lgpd_risk": v["lgpd_risk"]
        } for k, v in TAG_PATTERNS.items()}
    })


# ============================================================
# EXPORT ENDPOINTS - PDF e Excel
# ============================================================

@app.post("/export/pdf")
async def export_pdf(request: Request):
    """Gera relatório PDF profissional da auditoria"""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.units import cm
        
        data = await request.json()
        audit_result = data.get("result", {})
        
        if not audit_result:
            return JSONResponse({"status": "error", "msg": "Dados de auditoria não fornecidos"}, status_code=400)
        
        # Criar PDF em memória
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2*cm, bottomMargin=2*cm)
        elements = []
        styles = getSampleStyleSheet()
        
        # Estilo personalizado
        title_style = ParagraphStyle('CustomTitle', parent=styles['Heading1'], fontSize=24, textColor=colors.HexColor('#1e40af'), spaceAfter=20)
        subtitle_style = ParagraphStyle('CustomSubtitle', parent=styles['Heading2'], fontSize=14, textColor=colors.HexColor('#475569'), spaceAfter=10)
        normal_style = ParagraphStyle('CustomNormal', parent=styles['Normal'], fontSize=10, spaceAfter=8)
        
        # Cabeçalho
        elements.append(Paragraph("DataTrust Audit - Relatório de Auditoria", title_style))
        elements.append(Paragraph(f"URL: {audit_result.get('url', 'N/A')}", subtitle_style))
        elements.append(Paragraph(f"Data: {audit_result.get('timestamp', datetime.now().isoformat())[:10]}", normal_style))
        elements.append(Spacer(1, 20))
        
        # Score
        score = audit_result.get('score', 0)
        score_color = colors.green if score >= 70 else colors.orange if score >= 40 else colors.red
        elements.append(Paragraph(f"<b>Score de Conformidade LGPD:</b> <font color='{score_color}'>{score}%</font>", styles['Heading2']))
        elements.append(Spacer(1, 15))
        
        # Resumo
        summary = audit_result.get('summary', {})
        elements.append(Paragraph("<b>Resumo da Auditoria</b>", styles['Heading2']))
        summary_data = [
            ["Total de Tags", str(summary.get('totalTags', 0))],
            ["CMP Detectado", "Sim" if summary.get('consentDetected') else "Não"],
            ["Consent Mode v2", "Sim" if summary.get('consentModeV2') else "Não"],
            ["Violações LGPD", str(summary.get('violationsCount', 0))],
            ["Multa Estimada", summary.get('estimatedFine', 'R$ 0')]
        ]
        summary_table = Table(summary_data, colWidths=[8*cm, 8*cm])
        summary_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f1f5f9')),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('PADDING', (0, 0), (-1, -1), 8),
        ]))
        elements.append(summary_table)
        elements.append(Spacer(1, 20))
        
        # Tags detectadas
        tags = audit_result.get('tags', [])
        if tags:
            elements.append(Paragraph("<b>Tags Detectadas</b>", styles['Heading2']))
            tag_data = [["Tag", "Tipo", "Risco LGPD", "Linha"]]
            for tag in tags[:15]:  # Limitar a 15 tags
                tag_data.append([
                    tag.get('name', 'N/A'),
                    tag.get('type', 'N/A'),
                    tag.get('lgpdRisk', 'N/A'),
                    str(tag.get('lineNumber', 'N/A'))
                ])
            tag_table = Table(tag_data, colWidths=[6*cm, 4*cm, 3*cm, 3*cm])
            tag_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e40af')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
                ('PADDING', (0, 0), (-1, -1), 6),
                ('FONTSIZE', (0, 0), (-1, -1), 9),
            ]))
            elements.append(tag_table)
            elements.append(Spacer(1, 20))
        
        # Violações LGPD
        violations = audit_result.get('lgpdAnalysis', {}).get('violations', [])
        if violations:
            elements.append(Paragraph("<b>Violações LGPD Detectadas</b>", styles['Heading2']))
            for v in violations[:10]:  # Limitar a 10 violações
                elements.append(Paragraph(f"• <b>{v.get('tag', 'N/A')}</b> - {v.get('violation', 'N/A')}", normal_style))
                elements.append(Paragraph(f"  Artigo: {v.get('article', 'N/A')} | Multa: {v.get('estimatedFine', 'N/A')}", 
                    ParagraphStyle('Small', parent=normal_style, fontSize=9, textColor=colors.HexColor('#64748b'))))
        
        # Rodapé
        elements.append(Spacer(1, 30))
        elements.append(Paragraph("Relatório gerado por DataTrust Audit v2.0", 
            ParagraphStyle('Footer', parent=normal_style, fontSize=8, textColor=colors.HexColor('#94a3b8'))))
        
        doc.build(elements)
        buffer.seek(0)
        
        filename = f"gtm-audit-{audit_result.get('url', 'unknown').replace('https://', '').replace('/', '_')[:30]}.pdf"
        
        return StreamingResponse(
            buffer,
            media_type="application/pdf",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
        
    except Exception as e:
        print(f"❌ Erro ao gerar PDF: {e}")
        return JSONResponse({"status": "error", "msg": str(e)}, status_code=500)


@app.post("/export/excel")
async def export_excel(request: Request):
    """Gera relatório Excel detalhado da auditoria"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        
        data = await request.json()
        audit_result = data.get("result", {})
        
        if not audit_result:
            return JSONResponse({"status": "error", "msg": "Dados de auditoria não fornecidos"}, status_code=400)
        
        wb = Workbook()
        
        # ===== Aba: Resumo =====
        ws_summary = wb.active
        ws_summary.title = "Resumo"
        
        # Estilos
        header_fill = PatternFill(start_color="1e40af", end_color="1e40af", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        
        # Título
        ws_summary['A1'] = "DataTrust Audit - Relatório de Auditoria"
        ws_summary['A1'].font = Font(size=18, bold=True, color="1e40af")
        ws_summary.merge_cells('A1:D1')
        
        ws_summary['A3'] = "URL:"
        ws_summary['B3'] = audit_result.get('url', 'N/A')
        ws_summary['A4'] = "Data:"
        ws_summary['B4'] = audit_result.get('timestamp', '')[:10]
        ws_summary['A5'] = "Score LGPD:"
        ws_summary['B5'] = f"{audit_result.get('score', 0)}%"
        
        summary = audit_result.get('summary', {})
        ws_summary['A7'] = "Total de Tags:"
        ws_summary['B7'] = summary.get('totalTags', 0)
        ws_summary['A8'] = "CMP Detectado:"
        ws_summary['B8'] = "Sim" if summary.get('consentDetected') else "Não"
        ws_summary['A9'] = "Consent Mode v2:"
        ws_summary['B9'] = "Sim" if summary.get('consentModeV2') else "Não"
        ws_summary['A10'] = "Total Violações:"
        ws_summary['B10'] = summary.get('violationsCount', 0)
        ws_summary['A11'] = "Multa Estimada:"
        ws_summary['B11'] = summary.get('estimatedFine', 'R$ 0')
        
        # ===== Aba: Tags =====
        ws_tags = wb.create_sheet("Tags Detectadas")
        headers = ["Nome", "Tipo", "Risco LGPD", "Linha", "ID Detectado", "Antes do Consentimento"]
        for col, header in enumerate(headers, 1):
            cell = ws_tags.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin_border
        
        tags = audit_result.get('tags', [])
        for row, tag in enumerate(tags, 2):
            ws_tags.cell(row=row, column=1, value=tag.get('name', '')).border = thin_border
            ws_tags.cell(row=row, column=2, value=tag.get('type', '')).border = thin_border
            ws_tags.cell(row=row, column=3, value=tag.get('lgpdRisk', '')).border = thin_border
            ws_tags.cell(row=row, column=4, value=tag.get('lineNumber', '')).border = thin_border
            ws_tags.cell(row=row, column=5, value=tag.get('tagId', '')).border = thin_border
            ws_tags.cell(row=row, column=6, value="Sim" if tag.get('isBeforeConsent') else "Não").border = thin_border
        
        # Ajustar largura das colunas
        for col in range(1, 7):
            ws_tags.column_dimensions[get_column_letter(col)].width = 20
        
        # ===== Aba: Violações =====
        ws_violations = wb.create_sheet("Violações LGPD")
        v_headers = ["Tag", "Violação", "Artigo LGPD", "Severidade", "Multa Estimada"]
        for col, header in enumerate(v_headers, 1):
            cell = ws_violations.cell(row=1, column=col, value=header)
            cell.fill = PatternFill(start_color="dc2626", end_color="dc2626", fill_type="solid")
            cell.font = header_font
            cell.border = thin_border
        
        violations = audit_result.get('lgpdAnalysis', {}).get('violations', [])
        for row, v in enumerate(violations, 2):
            ws_violations.cell(row=row, column=1, value=v.get('tag', '')).border = thin_border
            ws_violations.cell(row=row, column=2, value=v.get('violation', '')).border = thin_border
            ws_violations.cell(row=row, column=3, value=v.get('article', '')).border = thin_border
            ws_violations.cell(row=row, column=4, value=v.get('severity', '')).border = thin_border
            ws_violations.cell(row=row, column=5, value=v.get('estimatedFine', '')).border = thin_border
        
        for col in range(1, 6):
            ws_violations.column_dimensions[get_column_letter(col)].width = 25
        
        # Salvar em memória
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        filename = f"gtm-audit-{audit_result.get('url', 'unknown').replace('https://', '').replace('/', '_')[:30]}.xlsx"
        
        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
        
    except Exception as e:
        print(f"❌ Erro ao gerar Excel: {e}")
        return JSONResponse({"status": "error", "msg": str(e)}, status_code=500)


@app.post("/chat")
async def chat(request: Request):
    """Chat endpoint com Gemini AI como orquestrador inteligente"""
    try:
        data = await request.json()
        message = data.get("message", "")
        context = data.get("context", {})
        
        # Se Gemini está configurado, usa IA
        if gemini_client:
            response, source = await generate_gemini_response(message, context)
        else:
            # Fallback para respostas baseadas em regras
            response = generate_chat_response(message, context)
            source = "rules_based"
        
        return JSONResponse({
            "status": "ok",
            "response": response,
            "source": source
        })
    except Exception as e:
        print(f"❌ Erro no chat: {e}")
        return JSONResponse({
            "status": "error",
            "msg": str(e)
        }, status_code=500)


@app.post("/chat/lgpd-bot")
async def chat_lgpd_bot(request: Request):
    """LGPD Compliance Bot - Auditor especializado com framework estruturado de 4 passos"""
    try:
        data = await request.json()
        message = data.get("message", "")
        context = data.get("context", {})
        if gemini_client:
            response, source = await generate_lgpd_bot_response(message, context)
        else:
            response = (
                "**[ERRO]** API de IA nao configurada. "
                "Configure GOOGLE_API_KEY no .env para ativar o LGPD Compliance Bot."
            )
            source = "no_ai"
        return JSONResponse({"status": "ok", "response": response, "source": source})
    except Exception as e:
        print(f"Erro no LGPD Bot: {e}")
        return JSONResponse({"status": "error", "msg": str(e)}, status_code=500)


async def generate_lgpd_bot_response(message: str, context: Dict) -> tuple:
    """LGPD Compliance Bot - System prompt especializado, framework de 4 passos, zero alucinacao"""
    try:
        tags_list = context.get("tags", [])
        pii_findings = context.get("personalDataFindings", [])
        sensitive_findings = context.get("sensitiveDataFindings", [])
        tags_before_consent = [t for t in tags_list if t.get("isBeforeConsent")]
        third_party_pii = [f for f in pii_findings if f.get("firstOrThirdParty") == "third_party"]

        tags_summary = "\n".join([
            f"  - {t.get('name','?')} ({t.get('type','?')}) | "
            f"Risco: {t.get('lgpdRisk','?')} | "
            f"Antes consentimento: {'SIM' if t.get('isBeforeConsent') else 'nao'}"
            for t in tags_list[:20]
        ]) or "  Nenhuma tag detectada."

        pii_summary = "\n".join([
            f"  - Campo: {f.get('fieldName','?')} | Tipo: {f.get('categoryLabel','?')} | "
            f"Fonte: {f.get('source','?')} | Destino: {f.get('destinationDomain','proprio')} | "
            f"3rd party: {'SIM' if f.get('firstOrThirdParty')=='third_party' else 'nao'}"
            for f in pii_findings[:10]
        ]) or "  Nenhum dado pessoal (Art. 5 I) detectado."

        sensitive_summary = "\n".join([
            f"  - Campo: {f.get('fieldName','?')} | Categoria: {f.get('categoryLabel','?')} | "
            f"Art. LGPD: {f.get('lgpdArticle','?')} | Consentimento: {f.get('consentState','?')}"
            for f in sensitive_findings[:10]
        ]) or "  Nenhum dado pessoal sensivel (Art. 5 II) detectado."

        consent_tags_summary = "\n".join([
            f"  - {t.get('name','?')} ({t.get('type','?')})"
            for t in tags_before_consent[:15]
        ]) or "  Nenhuma tag disparando antes do consentimento."

        system_prompt = f"""Voce e o LGPD Compliance Bot, Auditor e Especialista Tecnico em LGPD especializado em governanca de dados, privacidade tecnica (GTM, Data Layer, Cookies, CMP) e conformidade com a Lei 13.709/2018 e resolucoes da ANPD.

REGRAS INEGOCIAVEIS:
- Zero prolixidade. Nao use rodeios ou explicacoes academicas genericas.
- Estrutura estrita: toda resposta segue o framework de 4 secoes abaixo.
- Diferencie Dado Pessoal (Art. 5, I) de Dado Pessoal Sensivel (Art. 5, II) explicitamente.
- Cite apenas artigos reais da Lei 13.709/2018 e resolucoes reais da ANPD.
- Se dados estiverem incompletos, aponte o parametro faltante antes de assumir diagnostico.

FORMATO OBRIGATORIO — use EXATAMENTE esta estrutura Markdown:

## [DIAGNOSTICO TECNICO]
**Status Geral:** [Critico | Alto Risco | Medio Risco | Conforme]

**Infracoes Identificadas:**
| Elemento | Problema | Artigo Violado |
|---|---|---|
| (tag/campo/cookie) | (descricao tecnica) | Lei 13.709/2018 Art. XX |

## [IMPACTO REGULATORIO & RISCO]
- **Classificacao ANPD:** Infracao [Leve / Media / Grave]
- **Exposicao Financeira (Art. 52):** Ate 2% do faturamento, limite R$ 50 milhoes/infracao
- **Impacto Operacional:** (bloqueio de campanhas, reputacao, notificacoes)

## [PLANO DE REMEDIACAO TECNICO]
1. **Acao Imediata:** O que bloquear/pausar no GTM agora
2. **Correcao Estrutural:** Como sanitizar Data Layer, configurar CMP, hash SHA-256 de PII
3. **Validacao:** Como testar via console/preview do GTM

## [DISCLAIMER]
> Analise baseada em evidencias tecnicas do scanner DataTrust Audit. Nao substitui assessoria juridica. Consulte DPO ou advogado especialista em LGPD.

CONTEXTO DA AUDITORIA:
URL: {context.get('url', 'Nao informada')}
Score LGPD: {context.get('score', 0)}%
Violacoes: {context.get('violations', 0)}
CMP: {'DETECTADO' if context.get('hasConsentTool') else 'AUSENTE'}
Multa estimada: {context.get('estimatedFine', 'Nao calculada')}
Nivel exposicao: {context.get('exposureLevel', 'Nao avaliado')}

EVIDENCIAS TECNICAS COLETADAS:

[1] DISPARO PRECOCE — Tags antes do consentimento ({len(tags_before_consent)} detectadas):
{consent_tags_summary}

[2] PII — Dado Pessoal (Art. 5 I) — {len(pii_findings)} sinais, {len(third_party_pii)} para terceiros:
{pii_summary}

[3] DADO PESSOAL SENSIVEL (Art. 5 II) — {len(sensitive_findings)} sinais:
{sensitive_summary}

[4] TODAS AS TAGS ({len(tags_list)}):
{tags_summary}

Responda em Portugues Brasil com Markdown rico, tabelas e blocos de codigo quando aplicavel."""

        response = gemini_client.models.generate_content(
            model=GEMINI_MODEL,
            contents=message,
            config=types.GenerateContentConfig(
                system_instruction=system_prompt,
                temperature=0.25,
                max_output_tokens=3000
            )
        )
        return response.text, "lgpd_compliance_bot"

    except Exception as e:
        print(f"Erro no LGPD Bot Gemini: {e}")
        return (
            f"**[ERRO TECNICO]** Nao foi possivel processar: {str(e)}\n\n"
            "Verifique se GOOGLE_API_KEY esta configurada corretamente.",
            "error_fallback"
        )


async def generate_gemini_response(message: str, context: Dict) -> tuple:
    """Gera resposta usando Gemini AI como agente programador orquestrador"""
    try:
        # System prompt SUPER PODEROSO para agente programador
        system_prompt = f"""Você é o **DataTrust Audit AI**, um consultor sênior de Tech e Compliance.
    Você está analisando o site: {context.get('url', 'Desconhecido')}
    
    **Métricas Atuais:**
    - Score LGPD: {context.get('score', 0)}%
    - Violações: {context.get('violations', 0)}
    - CMP (Consentimento): {'Detectado' if context.get('hasConsentTool') else 'Ausente'}
    - Tags Encontradas: {len(context.get('tags', []))}
    - Multa Estimada: {context.get('estimatedFine', 'Indisponível')}
    - Vazamentos PII: {context.get('piiExposure', 0)}
    - Falhas Segurança: {context.get('securityIssues', 0)}
    
    **Diretrizes:**
    1. **Especialista Técnico**: Forneça trechos de código REAIS para corrigir as violações (ex: scripts para CMP, dataLayer.push seguro).
    2. **Especialista Jurídico**: Cite artigos específicos da LGPD (ex: Art. 7, Art. 18) ao explicar riscos.
    3. **Especialista em Performance**: Sugira melhorias na ordem de carregamento das tags.
    4. **Educador**: Explique o PORQUÊ de cada risco de forma clara mas profissional.
    
    **Poderes do DataTrust Audit AI:**
    - Você pode analisar tanto auditorias de página única quanto o **Deep Scan** (multi-página).
    - Se o usuário perguntar sobre "como corrigir", foque em soluções práticas de Tag Management.
    - Se detectar vazamento de PII (e-mails no dataLayer), priorize isso como ALERTA MÁXIMO.
    
    Responda em Português Brasil. Use Markdown rico com tabelas e blocos de código."""
        
        # Usa a nova API do Gemini
        response = gemini_client.models.generate_content(
            model=GEMINI_MODEL,
            contents=message,
            config=types.GenerateContentConfig(
                system_instruction=system_prompt,
                temperature=0.7,
                max_output_tokens=2048
            )
        )
        
        return response.text, "gemini_programmer_agent"
        
    except Exception as e:
        print(f"⚠️ Erro no Gemini, usando fallback: {e}")
        # Fallback para regras se Gemini falhar
        return generate_chat_response(message, context), "rules_fallback"


def generate_chat_response(message: str, context: Dict) -> str:
    """Gera resposta baseada em regras para o chat"""
    message_lower = message.lower()
    
    if "lgpd" in message_lower and "multa" in message_lower:
        return """📋 **Multas LGPD (Lei 13.709/2018)**

As multas por violação da LGPD podem chegar a:
- **2% do faturamento** da empresa, limitado a **R$ 50 milhões por infração**
- Multas diárias para forçar cumprimento
- Bloqueio ou eliminação dos dados pessoais

**Fatores agravantes:**
- Coleta sem consentimento
- Compartilhamento com terceiros sem base legal
- Uso de dados sensíveis (saúde, religião, etc.)
- Reincidência

🔗 Referência: Art. 52 da Lei 13.709/2018"""

    elif "consentimento" in message_lower or "cmp" in message_lower:
        return """🔐 **Consentimento e CMPs**

Para estar em conformidade com a LGPD, você precisa:

1. **Implementar uma CMP** (Consent Management Platform)
   - Cookiebot, OneTrust, ou Didomi são opções populares

2. **Carregar a CMP antes de qualquer tracking**
   - Deve ser o primeiro script no `<head>`

3. **Bloquear tags até o consentimento**
   - Configurar GTM com trigger de consentimento
   - Usar Consent Mode do Google

4. **Oferecer opções granulares**
   - Usuário deve poder aceitar/rejeitar categorias específicas

🔗 Referência: Art. 7º e 8º da LGPD"""

    elif "gtm" in message_lower or "tag manager" in message_lower:
        return """📊 **Google Tag Manager e LGPD**

O GTM em si **não coleta dados pessoais**, mas gerencia tags que coletam. Para conformidade:

1. **Configurar Consent Mode**
   - Habilitar integração com CMP
   - Definir estados de consentimento por categoria

2. **Criar triggers de consentimento**
   - Tags só devem disparar após consent

3. **Organizar por categoria**
   - Analytics, Marketing, Funcional, etc.

4. **Documentar Data Layer**
   - Mapear todos os dados que passam pelo GTM

⚠️ **Atenção:** Tags hardcoded no HTML que disparam antes do GTM violam LGPD!"""

    elif "universal analytics" in message_lower or " ua " in message_lower:
        return """⚠️ **Universal Analytics (Obsoleto)**

O Universal Analytics foi **descontinuado em julho de 2023**. Se ainda está no seu site:

❌ **Problemas:**
- Não recebe mais dados
- Código legado sem manutenção
- Potencial risco de segurança
- Não suporta Consent Mode moderno

✅ **O que fazer:**
1. Remover completamente UA do site e GTM
2. Migrar para Google Analytics 4
3. Configurar modo de consentimento
4. Atualizar políticas de privacidade

🔗 Referência: Art. 46 LGPD (medidas de segurança)"""

    else:
        violations = context.get("violations", 0)
        if violations > 0:
            return f"""📊 **Análise do seu site**

Foram detectadas **{violations} violações** de LGPD.

Principais recomendações:
1. Implementar uma ferramenta de consentimento (CMP)
2. Mover todas as tags de tracking para depois do consentimento
3. Configurar Google Tag Manager com Consent Mode
4. Revisar política de privacidade

Posso explicar mais sobre algum ponto específico?

💡 Pergunte sobre: multas, consentimento, GTM, ou qualquer tag específica."""
        else:
            return """👋 **Olá! Sou o assistente do DataTrust Audit.**

Posso ajudar com:
- 📋 Informações sobre multas LGPD
- 🔐 Como implementar consentimento
- 📊 Configuração de GTM
- ⚠️ Análise de violações

O que você gostaria de saber?"""


# SPA fallback must stay LAST so it does not shadow GET API/static routes.
@app.get("/{full_path:path}")
async def serve_spa(full_path: str):
    """Fallback route for client-side SPA navigation."""
    if full_path.startswith("assets/") or full_path.startswith("api/") or full_path.startswith("docs"):
        return JSONResponse({"error": "not found"}, status_code=404)
    index_path = FRONTEND_INDEX
    if not os.path.isfile(index_path):
        return JSONResponse(
            {"error": "frontend build not found", "health": "/health"},
            status_code=404,
        )
    return FileResponse(
        index_path,
        headers={"Cache-Control": "no-cache, no-store, must-revalidate", "Pragma": "no-cache"},
    )


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
