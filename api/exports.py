import io
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from datetime import datetime
import openpyxl

from db.database import get_db
from db.models import Scan, User, Subscription
from api.auth import get_current_user

router = APIRouter(prefix="/api/exports", tags=["Exports"])

@router.get("/xlsx/{scan_id}")
def export_scan_xlsx(scan_id: int, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    """Exports Scan findings to XLSX. Available only for Pro/Premium."""
    # Enforce Plan validation
    sub = db.query(Subscription).filter(Subscription.user_id == current_user.id).first()
    plan = sub.plan_name if sub else "free"
    
    if plan not in ["pro", "premium"]:
        raise HTTPException(status_code=403, detail="Exporte para XLSX está disponível apenas nos planos Pro e Premium.")
        
    scan = db.query(Scan).filter(Scan.id == scan_id, Scan.user_id == current_user.id).first()
    if not scan or not scan.raw_data:
        raise HTTPException(status_code=404, detail="Scan data not found or still processing.")
        
    wb = openpyxl.Workbook()
    
    # Overview Sheet
    ws1 = wb.active
    ws1.title = "Resumo do Scan"
    ws1.append(["URL Auditada:", scan.url])
    ws1.append(["Data do Scan:", str(scan.completed_at)])
    ws1.append(["Score de Conformidade:", f"{scan.score}/100"])
    
    data = scan.raw_data
    tags = data.get("tags", [])
    privacy = data.get("privacy", {})
    duplicates = data.get("duplicates", [])
    
    ws1.append(["Total de Tags Detectadas:", len(tags)])
    ws1.append(["Ferramenta de Consentimento:", "Sim" if privacy.get("has_consent_tool") else "Não"])
    if privacy.get("total_violations"):
        ws1.append(["Violações de Privacidade:", privacy.get("total_violations")])
        ws1.append(["Exposição Estimada:", privacy.get("estimatedRiskExposure")])
        
    # Tags Sheet
    ws2 = wb.create_sheet("Tags e Analytics")
    ws2.append(["ID", "Name", "Type", "Risk", "Found on Line", "Loaded Before Consent?"])
    for t in tags:
        ws2.append([
            t.get("tagId", "N/A"),
            t.get("name", "Unknown"),
            t.get("type", "Unknown"),
            t.get("lgpdRisk", "medium"),
            t.get("lineNumber", "Unknown"),
            "Sim" if t.get("isBeforeConsent") else "Não"
        ])
        
    # Privacy Violations Sheet
    if privacy.get("violations"):
        ws3 = wb.create_sheet("Violações de Privacidade")
        ws3.append(["Severidade", "Resumo", "Descrição", "Recomendação"])
        for v in privacy.get("violations"):
            ws3.append([
                v.get("severity", "medium"),
                v.get("summary", ""),
                v.get("description", ""),
                v.get("recommendation", "")
            ])
            
    # Duplicates Sheet
    if duplicates:
        ws4 = wb.create_sheet("Duplicidades Identificadas")
        ws4.append(["Severidade", "Entidade", "Nome", "Contagem", "Recomendação"])
        for d in duplicates:
            ws4.append([
                d.get("severity", ""),
                d.get("entity_type", ""),
                d.get("name", ""),
                d.get("occurrence_count", 0),
                d.get("recommendation", "")
            ])
    
    # Save to memory stream
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    
    filename = f"GTM_Audit_{urlparse(scan.url).netloc}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.get("/pdf/{scan_id}")
def export_scan_pdf(scan_id: int, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    """Exports Scan findings to PDF. Available only for Pro/Premium."""
    # Implement PDF generation logic using ReportLab 
    sub = db.query(Subscription).filter(Subscription.user_id == current_user.id).first()
    plan = sub.plan_name if sub else "free"
    
    if plan not in ["pro", "premium"]:
        raise HTTPException(status_code=403, detail="Exporte para PDF está disponível apenas nos planos Pro e Premium.")
        
    scan = db.query(Scan).filter(Scan.id == scan_id, Scan.user_id == current_user.id).first()
    if not scan or not scan.raw_data:
        raise HTTPException(status_code=404, detail="Scan data not found or still processing.")
        
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    
    stream = io.BytesIO()
    c = canvas.Canvas(stream, pagesize=letter)
    
    # Title
    c.setFont("Helvetica-Bold", 20)
    c.drawString(50, 750, "GTM Audit Pro - Relatório de Scan")
    
    c.setFont("Helvetica", 12)
    c.drawString(50, 720, f"URL: {scan.url}")
    c.drawString(50, 700, f"Data: {scan.completed_at}")
    c.drawString(50, 680, f"Score de Conformidade: {scan.score}/100")
    
    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, 640, "Resumo Executivo")
    c.setFont("Helvetica", 12)
    
    data = scan.raw_data
    privacy = data.get("privacy", {})
    tags = data.get("tags", [])
    duplicates = data.get("duplicates", [])
    
    c.drawString(50, 620, f"Total de Tags: {len(tags)}")
    has_cmp = "Sim" if privacy.get("has_consent_tool") else "Não"
    c.drawString(50, 600, f"Gestão de Consentimento Ativa (CMP): {has_cmp}")
    c.drawString(50, 580, f"Violações Encontradas: {privacy.get('total_violations', 0)}")
    
    if privacy.get("estimatedRiskExposure"):
         c.drawString(50, 560, f"Risco de Exposição: {privacy.get('estimatedRiskExposure')}")
         
    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, 520, "Tags Detectadas")
    c.setFont("Helvetica", 10)
    
    y = 500
    for idx, t in enumerate(tags[:10]):  # Limit to top 10 for basic PDF to avoid pagination complexity 
        tag_str = f"- {t.get('name')} ({t.get('type')}) - Risco LGPD: {t.get('lgpdRisk')} - Antes do Cookie: {'Sim' if t.get('isBeforeConsent') else 'Não'}"
        c.drawString(50, y, tag_str)
        y -= 20
        if y < 50:
             break
             
    if len(tags) > 10:
         c.drawString(50, y, f"... mais {len(tags) - 10} tags encontradas.")

    y -= 40
    if y > 100 and duplicates:
        c.setFont("Helvetica-Bold", 14)
        c.drawString(50, y, "Duplicidades Recorrentes")
        y -= 20
        c.setFont("Helvetica", 10)
        for d in duplicates[:5]:
            c.drawString(50, y, f"- {d.get('name')} (Risco: {d.get('severity')}) - {d.get('occurrence_count')}x")
            y -= 20
        
    c.save()
    stream.seek(0)
    
    from urllib.parse import urlparse
    filename = f"GTM_Audit_{urlparse(scan.url).netloc}.pdf"
    
    return StreamingResponse(
        stream,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
